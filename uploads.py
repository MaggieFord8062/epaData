"""
uploads.py
Reads an uploaded CSV or Excel file, checks it against the project schema,
and imports the approved records.

The steps follow the assignment's upload process:
  1. check the file extension and size          (check_file)
  2. read the file and find its columns         (read_rows, match_columns)
  3. compare the columns with the schema        (match_columns)
  4. find missing, invalid and duplicate rows   (analyze)
  5. show a preview and quality report          (the upload review page)
  6. store the approved records                 (import_upload)
Nothing is written to the emissions tables until the user approves.
"""
import csv
import json
import os
import re
from collections import Counter
from datetime import datetime

from models import Dataset, Facility, Unit, AnnualRecord
from validation import validate_record

UPLOAD_DIR = "uploads"
ALLOWED_EXTENSIONS = {"csv", "xlsx"}
MAX_FILE_BYTES = 10 * 1024 * 1024   # 10 MB
MAX_ROWS = 100_000

REQUIRED_FIELDS = ["facility_id", "facility_name", "state", "unit_id", "year"]

# Our field name -> column names a file might use for it: this site's CSV downloads,
# CAMPD's own column names, and plain snake_case
COLUMN_ALIASES = {
    "facility_id": ["Facility ID", "facilityId", "facility_id", "epa_facility_id", "ORISPL", "ORIS Code"],
    "facility_name": ["Facility", "Facility Name", "facilityName", "facility_name"],
    "state": ["State", "stateCode", "state_code", "state"],
    "county": ["County", "county"],
    "unit_id": ["Unit", "Unit ID", "unitId", "unit_id", "epa_unit_id"],
    "year": ["Year", "Reporting Year", "year", "reporting_year"],
    "unit_type": ["Unit Type", "unitType", "unit_type"],
    "primary_fuel": ["Primary Fuel", "Primary Fuel Type", "primaryFuelInfo", "primary_fuel"],
    "secondary_fuel": ["Secondary Fuel", "Secondary Fuel Type", "secondaryFuelInfo", "secondary_fuel"],
    "operating_time": ["Operating Time (hrs)", "Operating Time", "sumOpTime", "operating_time"],
    "gross_load": ["Gross Load (MWh)", "Gross Load", "grossLoad", "gross_load"],
    "steam_load": ["Steam Load (1000 lb)", "Steam Load", "steamLoad", "steam_load"],
    "heat_input": ["Heat Input (mmBtu)", "Heat Input", "heatInput", "heat_input"],
    "co2_mass": ["CO2 (tons)", "CO2 Mass", "CO2 Mass (short tons)", "co2Mass", "co2_mass"],
    "so2_mass": ["SO2 (tons)", "SO2 Mass", "SO2 Mass (short tons)", "so2Mass", "so2_mass"],
    "nox_mass": ["NOx (tons)", "NOx Mass", "NOx Mass (short tons)", "noxMass", "nox_mass"],
    "so2_control_info": ["SO2 Controls", "so2ControlInfo", "so2_control_info"],
    "nox_control_info": ["NOx Controls", "noxControlInfo", "nox_control_info"],
    "pm_control_info": ["PM Controls", "pmControlInfo", "pm_control_info"],
    "program_code": ["Programs", "Program Code", "programCodeInfo", "program_code"],
}

FIELD_LABELS = {
    "facility_id": "Facility ID", "facility_name": "Facility name", "state": "State", "county": "County",
    "unit_id": "Unit ID", "year": "Year", "unit_type": "Unit type", "primary_fuel": "Primary fuel",
    "secondary_fuel": "Secondary fuel", "operating_time": "Operating time", "gross_load": "Gross load",
    "steam_load": "Steam load", "heat_input": "Heat input", "co2_mass": "CO2", "so2_mass": "SO2",
    "nox_mass": "NOx", "so2_control_info": "SO2 controls", "nox_control_info": "NOx controls",
    "pm_control_info": "PM controls", "program_code": "Programs",
}


def normalize(name):
    """'CO2 (tons)', 'co2Mass' and 'co2_mass' become 'co2tons', 'co2mass', 'co2mass'."""
    return re.sub(r"[^a-z0-9]", "", str(name).lower())


ALIAS_LOOKUP = {normalize(alias): field for field, aliases in COLUMN_ALIASES.items() for alias in aliases}


# ---------- Step 1: extension and size ----------

def file_extension(filename):
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def check_file(filename, size):
    """Returns a list of problems with the file itself (empty if it's fine)."""
    problems = []
    extension = file_extension(filename)
    if not filename:
        problems.append("Choose a file to upload.")
    elif extension == "xls":
        problems.append("Old .xls Excel files aren't supported. Open it in Excel and save it as .xlsx, or export it as CSV.")
    elif extension not in ALLOWED_EXTENSIONS:
        kind = f".{extension} files" if extension else "Files without an extension"
        problems.append(f"{kind} aren't supported. Upload a CSV (.csv) or Excel (.xlsx) file.")
    if size == 0:
        problems.append("The file is empty.")
    elif size > MAX_FILE_BYTES:
        problems.append(f"The file is {size / 1024 / 1024:.1f} MB. The limit is {MAX_FILE_BYTES // 1024 // 1024} MB.")
    return problems


# ---------- Step 2: read the file ----------

def read_rows(path, extension):
    """Returns (column names, list of rows as {column: text})."""
    if extension == "csv":
        for encoding in ("utf-8-sig", "latin-1"):
            try:
                with open(path, newline="", encoding=encoding) as handle:
                    reader = csv.DictReader(handle)
                    rows = [row for _, row in zip(range(MAX_ROWS + 1), reader)]
                    return list(reader.fieldnames or []), rows
            except UnicodeDecodeError:
                continue
        raise ValueError("The CSV file's text encoding couldn't be read. Save it as UTF-8 and try again.")

    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    lines = sheet.iter_rows(values_only=True)
    header = next(lines, None) or []
    columns = [str(value).strip() if value is not None else "" for value in header]
    rows = []
    for values in lines:
        if len(rows) > MAX_ROWS:
            break
        if values is None or all(value in (None, "") for value in values):
            continue
        row = {}
        for column, value in zip(columns, values):
            if column:
                row[column] = "" if value is None else excel_text(value)
        rows.append(row)
    workbook.close()
    return [column for column in columns if column], rows


def excel_text(value):
    """Excel stores 2024 as 2024.0 and IDs like 1 as 1.0; turn whole numbers back into plain text."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


# ---------- Step 3: compare columns with the schema ----------

def match_columns(columns):
    matched = {}   # field -> file column
    ignored = []
    for column in columns:
        field = ALIAS_LOOKUP.get(normalize(column))
        if field and field not in matched:
            matched[field] = column
        else:
            ignored.append(column)
    return {
        "matched": [[FIELD_LABELS[field], column] for field, column in matched.items()],
        "fields": matched,
        "ignored": ignored,
        "missing_required": [FIELD_LABELS[field] for field in REQUIRED_FIELDS if field not in matched],
        "missing_optional": [FIELD_LABELS[field] for field in COLUMN_ALIASES
                             if field not in matched and field not in REQUIRED_FIELDS],
    }


# ---------- Step 4: find missing, invalid and duplicate rows ----------

def stored_keys(session):
    """(facility ID, unit ID, year) for every record already in the database."""
    rows = (
        session.query(AnnualRecord.epa_facility_id, Unit.epa_unit_id, AnnualRecord.year)
        .join(Unit, Unit.internal_unit_key == AnnualRecord.internal_unit_key)
        .all()
    )
    return {(facility_id, str(unit_id), year) for facility_id, unit_id, year in rows}


def warnings_for(record):
    """Rows that pass validation but look questionable. They're imported and listed in the report."""
    warnings = []
    emissions = [record.get(field) for field in ("co2_mass", "so2_mass", "nox_mass", "heat_input")]
    if all(value is None for value in emissions):
        warnings.append("No emissions or heat input values")
    if record.get("operating_time") == 0 and (record.get("co2_mass") or 0) > 0:
        warnings.append("Emissions reported with zero operating time")
    return warnings


def analyze(session, path, original_filename):
    """
    Read and check an uploaded file. Returns a summary that can be saved as JSON,
    shown on the review page, and used later to import the approved rows.
    """
    extension = file_extension(original_filename)
    columns, rows = read_rows(path, extension)
    columns_report = match_columns(columns)

    result = {
        "filename": original_filename,
        "columns": columns_report,
        "rows_read": len(rows),
        "too_many_rows": len(rows) > MAX_ROWS,
        "records": [],
        "problems": [],
        "missing_counts": {},
        "counts": {},
        "years": [],
    }
    if columns_report["missing_required"] or not rows or result["too_many_rows"]:
        return result

    fields = columns_report["fields"]
    already_stored = stored_keys(session)
    seen = {}
    counts = Counter()
    missing = Counter()

    for index, row in enumerate(rows):
        line = index + 2   # line 1 is the header
        record = {field: (row.get(column) or "").strip() for field, column in fields.items()}

        for field, value in record.items():
            if value == "":
                missing[FIELD_LABELS[field]] += 1

        # Excel sometimes turns unit IDs like 1 into 1.0
        if re.fullmatch(r"\d+\.0", record.get("unit_id", "")):
            record["unit_id"] = record["unit_id"][:-2]

        where = {"line": line, "facility_id": record.get("facility_id"), "unit_id": record.get("unit_id"),
                 "year": record.get("year")}

        clean, problem = validate_record(record)
        if problem:
            counts["rejected"] += 1
            result["problems"].append(dict(where, kind="Rejected", issue=problem))
            continue

        key = (clean["facility_id"], clean["unit_id"], clean["year"])
        if key in seen:
            counts["duplicate"] += 1
            result["problems"].append(dict(where, kind="Duplicate", issue=f"Same facility, unit and year as line {seen[key]}"))
            continue
        seen[key] = line

        if key in already_stored:
            counts["existing"] += 1
            result["problems"].append(dict(where, kind="Already stored", issue="This facility, unit and year is already in the database, so it will be skipped"))
            continue

        clean["county"] = (record.get("county") or "").strip() or None
        warnings = warnings_for(clean)
        if warnings:
            counts["questionable"] += 1
            for warning in warnings:
                result["problems"].append(dict(where, kind="Questionable", issue=warning))

        counts["ready"] += 1
        result["records"].append({"line": line, "record": clean})

    result["counts"] = {
        "ready": counts["ready"],
        "questionable": counts["questionable"],
        "rejected": counts["rejected"],
        "duplicate": counts["duplicate"],
        "existing": counts["existing"],
    }
    result["missing_counts"] = dict(missing.most_common())
    result["years"] = sorted({item["record"]["year"] for item in result["records"]})
    return result


def save_analysis(upload_id, result):
    with open(analysis_path(upload_id), "w") as handle:
        json.dump(result, handle)


def load_analysis(upload_id):
    try:
        with open(analysis_path(upload_id)) as handle:
            return json.load(handle)
    except FileNotFoundError:
        return None


def analysis_path(upload_id):
    return os.path.join(UPLOAD_DIR, f"{upload_id}_analysis.json")


def quality_report_rows(result):
    """The data-quality report as rows for a CSV download."""
    return [
        {
            "Line": problem["line"],
            "Status": problem["kind"],
            "Facility ID": problem["facility_id"],
            "Unit ID": problem["unit_id"],
            "Year": problem["year"],
            "Issue": problem["issue"],
        }
        for problem in result.get("problems", [])
    ]


# ---------- Step 6: store the approved records ----------

def year_label(years):
    if not years:
        return None
    if len(years) == 1:
        return str(years[0])
    return f"{years[0]}-{years[-1]}"


def import_upload(session, upload, result):
    """
    Store the rows that passed. Rows already in the database are skipped, never overwritten.
    Creates the Dataset (provenance) record and links every stored row to it.
    """
    records = result.get("records", [])
    already_stored = stored_keys(session)   # checked again in case something changed since the preview

    dataset = Dataset(
        dataset_name=f"Upload: {upload.original_filename}",
        data_source="User upload",
        reporting_year=year_label(result.get("years")),
        retrieval_date=datetime.now().isoformat(timespec="seconds"),
        og_filename=upload.original_filename,
        num_raw_records=result.get("rows_read", 0),
    )
    session.add(dataset)
    session.flush()

    facilities = {f.epa_facility_id: f for f in session.query(Facility).all()}
    units = {(u.epa_facility_id, u.epa_unit_id): u for u in session.query(Unit).all()}

    inserted = 0
    skipped_now = 0
    new_facilities = 0
    new_units = 0

    for item in records:
        record = item["record"]
        key = (record["facility_id"], record["unit_id"], record["year"])
        if key in already_stored:
            skipped_now += 1
            continue

        facility = facilities.get(record["facility_id"])
        if facility is None:
            facility = Facility(
                epa_facility_id=record["facility_id"],
                facility_name=record["facility_name"],
                state=record["state"],
                county=record.get("county"),
            )
            session.add(facility)
            session.flush()
            facilities[record["facility_id"]] = facility
            new_facilities += 1

        unit = units.get((record["facility_id"], record["unit_id"]))
        if unit is None:
            unit = Unit(
                epa_facility_id=record["facility_id"],
                epa_unit_id=record["unit_id"],
                unit_type=record.get("unit_type"),
                primary_fuel=record.get("primary_fuel"),
                secondary_fuel=record.get("secondary_fuel"),
            )
            session.add(unit)
            session.flush()
            units[(record["facility_id"], record["unit_id"])] = unit
            new_units += 1

        session.add(AnnualRecord(
            epa_facility_id=record["facility_id"],
            internal_unit_key=unit.internal_unit_key,
            year=record["year"],
            dataset_id=dataset.dataset_id,
            operating_time=record.get("operating_time"),
            gross_load=record.get("gross_load"),
            steam_load=record.get("steam_load"),
            heat_input=record.get("heat_input"),
            co2_mass=record.get("co2_mass"),
            so2_mass=record.get("so2_mass"),
            nox_mass=record.get("nox_mass"),
            so2_control_info=record.get("so2_control_info"),
            nox_control_info=record.get("nox_control_info"),
            pm_control_info=record.get("pm_control_info"),
            program_code=record.get("program_code"),
        ))
        already_stored.add(key)
        inserted += 1

    counts = result.get("counts", {})
    dataset.num_accepted_records = inserted

    notes = [f"{inserted} new records stored"]
    if new_facilities or new_units:
        notes.append(f"{new_facilities} new facilities and {new_units} new units added")
    if counts.get("questionable"):
        notes.append(f"{counts['questionable']} questionable rows stored and listed in the quality report")
    if counts.get("existing") or skipped_now:
        notes.append(f"{counts.get('existing', 0) + skipped_now} rows skipped: already in the database")
    if counts.get("duplicate"):
        notes.append(f"{counts['duplicate']} duplicate rows in the file skipped")
    if counts.get("rejected"):
        reasons = Counter(problem["issue"] for problem in result.get("problems", []) if problem["kind"] == "Rejected")
        for reason, count in reasons.most_common():
            notes.append(f"Rejected ({count}): {reason}")
    dataset.notes = "; ".join(notes)

    upload.status = "imported"
    upload.dataset_id = dataset.dataset_id
    session.commit()
    return dataset